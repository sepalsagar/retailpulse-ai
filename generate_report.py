"""Generate the RetailPulse AI internship report from validated local outputs."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data" / "processed"
SCREENSHOTS = ROOT / "assets" / "screenshots"


def gbp(value: float) -> str:
    return f"£{float(value):,.2f}"


def count(value: int | float) -> str:
    return f"{int(value):,}"


def percent(value: float, decimals: int = 1) -> str:
    return f"{float(value):.{decimals}%}"


def read_inputs() -> dict:
    required = [
        "kpi_summary.json", "cancellation_return_summary.json", "validation_report.json",
        "cluster_profiles.parquet", "clustering_evaluation.parquet", "product_metrics.parquet",
        "country_metrics.parquet", "opportunities.parquet", "recommendations.parquet",
        "customer_risk.parquet", "return_product_metrics.parquet",
    ]
    missing = [name for name in required if not (DATA / name).is_file()]
    expected_shots = [
        "01_executive_overview.png", "02_customer_intelligence.png",
        "03_product_market.png", "04_risks_opportunities.png",
    ]
    missing += [f"assets/screenshots/{name}" for name in expected_shots if not (SCREENSHOTS / name).is_file()]
    if missing:
        raise FileNotFoundError("Required report inputs are missing: " + ", ".join(missing))
    return {
        "kpi": json.loads((DATA / "kpi_summary.json").read_text(encoding="utf-8")),
        "returns": json.loads((DATA / "cancellation_return_summary.json").read_text(encoding="utf-8")),
        "validation": json.loads((DATA / "validation_report.json").read_text(encoding="utf-8")),
        "clusters": pd.read_parquet(DATA / "cluster_profiles.parquet").sort_values("cluster_id"),
        "evaluation": pd.read_parquet(DATA / "clustering_evaluation.parquet").sort_values("k"),
        "products": pd.read_parquet(DATA / "product_metrics.parquet"),
        "countries": pd.read_parquet(DATA / "country_metrics.parquet"),
        "opportunities": pd.read_parquet(DATA / "opportunities.parquet"),
        "recommendations": pd.read_parquet(DATA / "recommendations.parquet"),
        "risk": pd.read_parquet(DATA / "customer_risk.parquet"),
        "return_products": pd.read_parquet(DATA / "return_product_metrics.parquet"),
        "screenshots": expected_shots,
    }


def set_cell(cell, value: object, *, header: bool = False) -> None:
    cell.text = str(value)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(8)
            if header:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    if header:
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "17365D")
        cell._tc.get_or_add_tcPr().append(shade)


def add_table(doc: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell(cell, header, header=True)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            set_cell(cell, value)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def configure_section(section, landscape: bool = False) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width, section.page_height = Inches(11), Inches(8.5)
        section.left_margin = section.right_margin = Inches(.55)
        section.top_margin = section.bottom_margin = Inches(.5)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(.7)
        section.top_margin = section.bottom_margin = Inches(.6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "RetailPulse AI | Internship Project Report | "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build_report(data: dict) -> Path:
    k, r, v = data["kpi"], data["returns"], data["validation"]
    clusters, evaluation = data["clusters"], data["evaluation"]
    products, countries, risk = data["products"], data["countries"], data["risk"]
    doc = Document()
    configure_section(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9)
    for name, size, rgb in [("Title", 28, (20, 43, 69)), ("Heading 1", 15, (20, 43, 69)), ("Heading 2", 11, (36, 99, 166))]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*rgb)

    # 1. Title page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("RETAILPULSE AI")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(36, 99, 166)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Retail Revenue & Customer\nIntelligence System")
    for line in [
        "AICTE | IBM SkillsBuild Data Analytics with AI Internship 2026",
        "Student Project", "Sepal Sagar", "B.Tech CSE",
        "Bengal College of Engineering and Technology",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_heading(doc, "2. Abstract")
    add_paragraph(doc, f"RetailPulse AI converts {count(v['rows_raw'])} UCI Online Retail II transaction rows into validated KPIs, product and market analysis, RFM profiles, descriptive K-Means segments, a transparent rule-based risk indicator, and evidence-based opportunities. A four-page Streamlit dashboard presents the results for business review. It supports decision-making but does not predict churn or establish causal effects.")
    add_heading(doc, "3. Introduction")
    add_paragraph(doc, "Transaction records alone do not explain performance, its drivers, customer inactivity, or management choices. The project makes a traceable analytical path from operational records to business action.")
    add_paragraph(doc, "DATA  →  INFORMATION  →  INSIGHTS  →  DECISION  →  ACTION")
    add_heading(doc, "4. Problem Statement")
    add_paragraph(doc, "A retail business has a large volume of historical transaction data but needs a unified intelligence system to understand revenue performance, customer value, customer inactivity, product performance, market performance, risks, and opportunities. The system should help management move from descriptive metrics toward actionable decisions.")
    add_heading(doc, "5. Business Objectives")
    add_bullets(doc, ["Measure business performance with defined KPIs and trends.", "Identify important products, markets, and revenue concentration.", "Segment identified customers using RFM behavior.", "Prioritize valuable customers with elevated rule-based risk scores.", "Identify evidence-backed product, market, and customer opportunities.", "Translate findings into recommended actions through an interactive dashboard."])

    add_heading(doc, "6. Dataset")
    add_paragraph(doc, f"Source: UCI Machine Learning Repository, Online Retail II (dataset 502), credited to Chen (2012). The source workbook has {count(v['rows_raw'])} rows and eight fields: Invoice, StockCode, Description, Quantity, InvoiceDate, Price, Customer ID, and Country. It covers {v['date_min'][:10]} to {v['date_max'][:10]}. The approximately 43.5 MB workbook has two overlapping sheets combined using actual invoice dates. Currency is pound sterling (GBP).")
    add_heading(doc, "7. Data Quality")
    add_paragraph(doc, f"Invoice, StockCode, Quantity, InvoiceDate, Price, and Country have no missing values. Description is missing on {count(v['missing_values_by_column']['Description'])} rows; Customer ID is missing on {count(v['missing_values_by_column']['Customer ID'])} rows. There are {count(v['distinct_invoices'])} eligible invoices, {count(v['distinct_customers_on_sales'])} identified customers, {count(v['distinct_products_on_sales'])} products, and {count(v['distinct_countries_on_sales'])} countries. Before deduplication, {count(r['negative_quantity_without_c_prefix_lines_before_dedup'])} negative-quantity lines have no C prefix and all are zero-priced; {count(r['zero_price_positive_quantity_lines_before_dedup'])} positive-quantity zero-price lines represent {count(r['zero_price_positive_quantity_units_separate_before_dedup'])} units. Five negative-price bad-debt adjustments and one C-prefixed nonnegative-quantity anomaly were found.")

    add_heading(doc, "8. Data Cleaning")
    add_heading(doc, "Duplicate handling", 2)
    add_paragraph(doc, f"Within-sheet duplicate checks and cross-sheet 64-bit row fingerprints found {count(v['duplicate_occurrences_removed'])} repeated full-row occurrences across {count(v['exact_duplicate_pattern_count'])} patterns. One copy was retained in aggregates; the source workbook was unchanged. This removed {count(v['duplicate_removed_ordinary_sales_lines'])} eligible sale lines, {count(v['duplicate_removed_ordinary_sales_units'])} units, and {gbp(v['duplicate_removed_ordinary_sales_value'])} gross sales ({percent(v['duplicate_removed_ordinary_sales_value'] / k['gross_positive_sales_value'])} of final gross). The lines affected {count(v['duplicate_affected_ordinary_sales_invoices'])} invoices, but no eligible invoice was removed. The treatment materially affects revenue and is disclosed.")
    add_heading(doc, "Transaction classification", 2)
    add_bullets(doc, [
        "Sale: numeric invoice, no C prefix, merchandise classification, positive quantity and price, and not an exact duplicate.",
        "Defined merchandise return: C-prefixed invoice, negative quantity, positive price, merchandise classification, and not duplicate. Its signed value is negative and reported separately.",
        "Postage/carriage services and their qualifying reversals are reported separately. Gift vouchers, accounting entries, discounts, tests, commissions, and adjustments are not merchandise sales or returns.",
        f"Zero-price rows are excluded from paid sales; positive-quantity zero-price lines ({count(r['zero_price_positive_quantity_lines_before_dedup'])}) and negative-quantity zero-price signals ({count(r['negative_quantity_zero_price_lines_before_dedup'])}) remain separate. The five negative-price bad-debt rows total {gbp(r['negative_price_bad_debt_adjustment_value_before_dedup'])} before deduplication.",
        "Missing Customer IDs remain in business-wide aggregates but are excluded from customer-level analysis. Missing product descriptions receive a fallback label.",
    ])

    add_heading(doc, "9. KPI Framework")
    add_paragraph(doc, "Headline sales are merchandise-only; service charges and reversals are reported separately.")
    add_table(doc, ["KPI", "Baseline", "Definition"], [
        ["Gross Sales", gbp(k["gross_positive_sales_value"]), "Eligible positive Quantity × Price merchandise lines after deduplication."],
        ["Return value", gbp(k["signed_cancellation_return_value"]), "Signed value of defined C-prefixed merchandise returns."],
        ["Net Sales", gbp(k["net_sales_after_defined_returns"]), "Gross Sales plus signed defined returns."],
        ["Orders", count(k["orders"]), "Distinct invoices with eligible sale lines."],
        ["Identified Customers", count(k["customers"]), "Distinct nonmissing IDs on eligible sales."],
        ["Paid Units", count(k["units"]), "Eligible positive-quantity, positive-price merchandise units."],
        ["AOV", gbp(k["average_order_value"]), "Gross Sales divided by eligible orders."],
        ["Repeat Customer Rate", percent(k["repeat_customer_rate"], 2), "Identified customers with >1 eligible orders divided by identified customers."],
        ["Return Line Rate", percent(k["c_invoice_return_line_rate"], 3), "Defined return lines divided by eligible sale lines."],
        ["Comparable Growth", percent(k["revenue_growth_2011_jan_nov_vs_2010_jan_nov"], 2), "Jan–Nov 2011 gross divided by Jan–Nov 2010 gross, minus one."],
    ])
    add_paragraph(doc, f"Service charges are {gbp(k['service_charge_revenue_separate'])}; service reversals are {gbp(k['service_charge_cancellation_value_separate'])}.")

    add_heading(doc, "10. Exploratory Analysis")
    add_paragraph(doc, f"Comparable Jan–Nov sales rose from {gbp(k['growth_revenue_2010_jan_nov'])} to {gbp(k['growth_revenue_2011_jan_nov'])}, or {percent(k['revenue_growth_2011_jan_nov_vs_2010_jan_nov'], 2)}. December is excluded because 2011 ends on December 9. The top 10 identified customers contribute {percent(k['top_10_identified_customer_revenue_share'], 2)} of identified revenue; the top 10 products contribute {percent(k['top_10_merchandise_product_revenue_share'], 2)} of merchandise sales. Defined merchandise returns are {count(r['c_prefix_negative_quantity_price_positive_lines_after_dedup'])} lines, {count(abs(r['c_prefix_negative_quantity_signed_units_after_dedup']))} units, and {gbp(r['c_prefix_negative_quantity_signed_value_after_dedup'])} signed value.")
    add_heading(doc, "11. Customer Intelligence")
    add_paragraph(doc, f"RFM uses identified customers and eligible merchandise sales. Recency is days since last eligible purchase, Frequency is distinct eligible invoices, and Monetary is eligible gross merchandise sales. The reference date is 2011-12-10, one day after the latest valid sale. {count(k['raw_transaction_lines_with_customer_id'])} source rows have Customer ID and {count(k['raw_transaction_lines_without_customer_id'])} do not. Identified customers represent {percent(k['identified_customer_revenue_share'])} ({gbp(k['identified_customer_revenue'])}) of eligible merchandise revenue.")

    add_heading(doc, "12. Customer Segmentation")
    add_paragraph(doc, "RFM features are log1p transformed and standardized before K-Means. K=2–6 were assessed with sampled silhouette (up to 2,500 customers) and adjusted Rand seed stability. K=2 has the highest silhouette, but K=4 was selected for actionable profile separation and strong stability. These are behavioral segments, not supervised predictions.")
    add_table(doc, ["K", "Sampled silhouette", "Seed stability ARI"], [[int(row.k), f"{row.silhouette_sampled:.3f}", f"{row.seed_stability_ari:.3f}"] for row in evaluation.itertuples()])
    add_table(doc, ["Segment", "Customers", "Avg monetary", "Avg orders", "Avg recency", "Revenue share"], [[row.segment, count(row.customers), gbp(row.avg_monetary), f"{row.avg_frequency:.2f}", f"{row.avg_recency_days:.0f} days", percent(row.revenue_share)] for row in clusters.itertuples()])
    add_paragraph(doc, "The four profiles are Low-Value Inactive, Recent Developing, Valuable At Risk, and High-Value Active. Names were assigned after profile review.")

    add_heading(doc, "13. Customer Risk")
    add_paragraph(doc, "The Customer Risk Indicator is a rule-based risk indicator, not a churn prediction model. Score = 100 × (0.50 × recency percentile + 0.30 × six-month activity decline + 0.20 × defined return-value ratio). Activity compares Jun–Nov 2011 spend with Dec 2010–May 2011. Bands: Low ≤50, Moderate >50–75, High >75. High-value at-risk status requires top-quartile monetary value and score ≥60.")
    at_risk = risk.loc[risk.is_high_value_at_risk]
    add_paragraph(doc, f"{count(len(at_risk))} customers meet the high-value at-risk rule, with {gbp(at_risk.customer_value.sum())} historical eligible sales. This is a prioritization signal, not predicted customer loss.")

    add_heading(doc, "14. Product Intelligence")
    top_products = products.sort_values("sales_revenue", ascending=False).head(5)
    add_table(doc, ["Product", "Gross sales", "Revenue share", "Orders", "Return line rate"], [[f"{row.product_description} ({row.stock_code})", gbp(row.sales_revenue), percent(row.revenue_share, 2), count(row.sales_orders), percent(row.return_line_rate, 2)] for row in top_products.itertuples()])
    opportunities = data["opportunities"]
    product_opps = opportunities.loc[opportunities.category.eq("Strong growing product")]
    if not product_opps.empty:
        add_paragraph(doc, "Examples of qualifying product opportunities from the implemented engine:")
        add_bullets(doc, [str(x) for x in product_opps.evidence.head(3)])
    unusual = v["high_volume_single_order_products"][0]
    add_paragraph(doc, f"An unusual-volume example is product {unusual['stock_code']} ({unusual['product_description']}): {count(unusual['sales_units'])} units in one order and an equal defined return quantity. The pattern warrants review and is not treated as repeatable demand.")
    return_products = data["return_products"].sort_values("returns_revenue")
    if not return_products.empty:
        top_return = return_products.iloc[0]
        add_paragraph(doc, f"The largest absolute product return value is for code {top_return.stock_code}: {gbp(top_return.returns_revenue)}. Return causes are not inferred from transaction data.")

    add_heading(doc, "15. Market Intelligence")
    top_countries = countries.sort_values("sales_revenue", ascending=False).head(5)
    add_table(doc, ["Market", "Gross sales", "Business share", "Customers", "Jan–Nov growth"], [[row.country, gbp(row.sales_revenue), percent(row.revenue_share), count(row.sales_customers), "n/a" if pd.isna(row.revenue_growth) else percent(row.revenue_growth)] for row in top_countries.itertuples()])
    uk = countries.loc[countries.country.eq("United Kingdom")].iloc[0]
    france = countries.loc[countries.country.eq("France")].iloc[0]
    add_paragraph(doc, f"The UK contributes {percent(uk.revenue_share)} of merchandise sales. This is exposure to the largest market, not evidence it is inherently undesirable. France grew {percent(france.revenue_growth)} with {count(france.orders_2010_jan_nov)} and {count(france.orders_2011_jan_nov)} orders in the comparable windows, satisfying the 100-order support rule.")

    add_heading(doc, "16. Risks")
    add_bullets(doc, [
        f"{count(len(at_risk))} high-value customers meet the rule-based risk threshold.",
        f"Defined merchandise reversals total {gbp(r['c_prefix_negative_quantity_signed_value_after_dedup'])}; unpriced signals are separate.",
        f"The UK represents {percent(uk.revenue_share)} of eligible merchandise sales.",
        f"{percent(k['unidentified_customer_revenue_share'])} of eligible revenue has no Customer ID.",
        f"Exact duplicate handling changes gross sales by {gbp(v['duplicate_removed_ordinary_sales_value'])}.",
        "Historical transactions do not establish causation or current conditions; the risk score is not churn prediction.",
    ])

    add_heading(doc, "17. Opportunities")
    for row in opportunities.itertuples():
        add_heading(doc, str(row.category), 2)
        add_paragraph(doc, "Evidence: " + str(row.evidence))
        add_paragraph(doc, "Interpretation: " + str(row.business_meaning))
        add_paragraph(doc, "Business impact: historical evidence identifies a review or test opportunity; realized impact is not measured.")
        add_paragraph(doc, "Recommended action: " + str(row.recommended_action))

    add_heading(doc, "18. Business Recommendations")
    recs = data["recommendations"].sort_values("priority", key=lambda s: s.map({"High": 0, "Medium": 1, "Low": 2}))
    for row in recs.itertuples():
        add_heading(doc, f"{row.priority} priority — {row.segment}", 2)
        add_paragraph(doc, f"{row.observation} {row.insight} {row.business_impact}")
        add_paragraph(doc, "Next action: " + str(row.recommended_action))
    add_paragraph(doc, "Prioritize a measured retention test for the valuable at-risk group, protect active customer experience, and validate stock, margin, returns, and customer repeat behavior before expansion. The project does not estimate realized uplift.")

    add_heading(doc, "19. System Architecture")
    add_paragraph(doc, "UCI Online Retail II → sheet-wise cleaning and classification → validated Parquet/JSON tables → KPI and trend analysis → RFM and K-Means → rule-based risk and opportunity engine → Streamlit dashboard → business decisions.")
    add_paragraph(doc, "The implementation does not use a database, external API, cloud infrastructure, or LLM.")
    add_heading(doc, "20. Technology Stack")
    add_bullets(doc, ["Python, pandas, NumPy, and openpyxl for analysis and workbook access.", "scikit-learn for scaling, K-Means, silhouette evaluation, and adjusted Rand stability.", "PyArrow for Parquet output; Streamlit and Plotly for the dashboard and charts.", "python-docx for report generation."])

    # Landscape pages preserve readability of the actual wide dashboard captures.
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape, landscape=True)
    add_heading(doc, "21. Dashboard Screenshots")
    figure_names = ["Executive Overview", "Customer Intelligence", "Product & Market Intelligence", "Risks & Opportunities"]
    for index, (name, filename) in enumerate(zip(figure_names, data["screenshots"]), 1):
        if index > 1:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SCREENSHOTS / filename), width=Inches(9.7))
        caption = doc.add_paragraph(f"Figure {index}. RetailPulse AI — {name}.")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True

    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(portrait)
    add_heading(doc, "22. Key Results")
    add_table(doc, ["Measure", "Validated result"], [
        ["Gross merchandise sales", gbp(k["gross_positive_sales_value"])],
        ["Net merchandise sales", gbp(k["net_sales_after_defined_returns"])],
        ["Orders", count(k["orders"])], ["Identified customers", count(k["customers"])],
        ["Average order value", gbp(k["average_order_value"])],
        ["Repeat customer rate", percent(k["repeat_customer_rate"], 2)],
        ["Comparable Jan–Nov growth", percent(k["revenue_growth_2011_jan_nov_vs_2010_jan_nov"], 2)],
        ["UK revenue share", percent(uk.revenue_share)],
        ["Valuable At Risk customers", count(clusters.loc[clusters.segment.eq("Valuable At Risk"), "customers"].iloc[0])],
        ["High-value at-risk opportunity", f"{count(len(at_risk))} customers / {gbp(at_risk.customer_value.sum())}"],
    ])
    add_heading(doc, "23. Limitations")
    add_bullets(doc, [
        f"{count(k['raw_transaction_lines_without_customer_id'])} source rows lack Customer ID; customer analysis covers identified transactions ({percent(k['identified_customer_revenue_share'])} of eligible revenue).",
        "Historical analysis is descriptive, does not establish causation, and may not represent current conditions.",
        "The risk indicator is rule-based, not churn prediction; clustering describes observed behavior.",
        "Duplicate treatment materially affects gross sales; unpriced quantity events do not support monetary return calculations.",
        "No margin, campaign exposure, or verified churn labels are available; opportunity values are historical, not forecast uplift.",
    ])
    add_heading(doc, "24. Future Scope")
    add_bullets(doc, ["Real-time ingestion and alerts.", "Predictive churn analysis only if reliable labeled outcomes become available.", "Demand forecasting and additional customer features.", "Database-backed architecture, cloud deployment, and a governed natural-language analytics assistant."])
    add_paragraph(doc, "These are future possibilities, not implemented features.")
    add_heading(doc, "25. Conclusion")
    add_paragraph(doc, f"RetailPulse AI converts {count(v['rows_raw'])} historical transaction rows into validated KPIs, trends, drivers, risk priorities, opportunities, and recommended actions: KPI → Trend → Driver → Risk → Opportunity → Action. It provides an auditable basis for management review without claims of prediction or causation.")
    add_heading(doc, "26. References")
    add_bullets(doc, [
        "Chen, D. (2012). Online Retail II [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5CG6D",
        "UCI Online Retail II dataset page: https://archive.ics.uci.edu/dataset/502/online%2Bretail%2Bii",
        "Streamlit documentation: https://docs.streamlit.io/",
        "scikit-learn documentation: https://scikit-learn.org/stable/",
        "Plotly Python documentation: https://plotly.com/python/",
        "pandas documentation: https://pandas.pydata.org/docs/",
        "AICTE | IBM SkillsBuild Data Analytics with AI Internship 2026 project context supplied for this project.",
    ])
    path = ROOT / "PROJECT_REPORT.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    output = build_report(read_inputs())
    print(f"Generated {output.name} ({output.stat().st_size:,} bytes)")
